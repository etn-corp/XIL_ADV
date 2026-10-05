#!/usr/bin/env python3
from __future__ import annotations
import argparse, re, sys
from dataclasses import dataclass, field
from pathlib import Path
from openpyxl import load_workbook

HEADERS=("Test Case Name","Description","Action","Expected Result")
EXTENSIONS={".xlsx",".xlsm"}
@dataclass
class Step: action:str=""; expected:str=""
@dataclass
class Case: name:str; descriptions:list[str]=field(default_factory=list); steps:list[Step]=field(default_factory=list)

def clean(v):
    if v is None:return ""
    lines=[x.rstrip() for x in str(v).replace("\r\n","\n").replace("\r","\n").split("\n")]
    while lines and not lines[0].strip():lines.pop(0)
    while lines and not lines[-1].strip():lines.pop()
    return "\n".join(lines)

def slug(v):
    v=re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',v.strip());v=re.sub(r'\s+','_',v);v=re.sub(r'_+','_',v).strip('._')
    return v or 'unnamed_test_case'

def header(ws):
    for row in ws.iter_rows():
        found={clean(c.value):c.column for c in row if clean(c.value)}
        if all(h in found for h in HEADERS):return row[0].row,{h:found[h] for h in HEADERS}
    return None

def parse(ws):
    h=header(ws)
    if not h:return []
    hr,col=h; out=[]; current=None
    for r in range(hr+1,ws.max_row+1):
        name=clean(ws.cell(r,col['Test Case Name']).value);desc=clean(ws.cell(r,col['Description']).value)
        action=clean(ws.cell(r,col['Action']).value);expected=clean(ws.cell(r,col['Expected Result']).value)
        if name:
            if current:out.append(current)
            current=Case(name)
        if not current:continue
        if desc and desc not in current.descriptions:current.descriptions.append(desc)
        if action or expected:current.steps.append(Step(action,expected))
    if current:out.append(current)
    return out

def cell(v,empty=''):
    return (v or empty).replace('\\','\\\\').replace('|','\\|').replace('\n','<br>')

def render(c):
    lines=['| Test case name | Description |','|---|---|',f"| {cell(c.name)} | {cell(chr(10).join(c.descriptions),'(not specified)')} |",'',
           '| Step | Action | Expected result |','|---:|---|---|']
    if not c.steps:lines.append('| 1 | (not specified) | (not specified) |')
    for i,s in enumerate(c.steps,1):lines.append(f"| {i} | {cell(s.action,'(not specified)')} | {cell(s.expected,'(not specified)')} |")
    return '\n'.join(lines)+'\n'

def inputs(items):
    result=set()
    for x in items:
        p=Path(x)
        if p.is_dir():result.update(q for q in p.rglob('*') if q.suffix.lower() in EXTENSIONS and not q.name.startswith('~$'))
        elif p.is_file() and p.suffix.lower() in EXTENSIONS and not p.name.startswith('~$'):result.add(p)
        else:print(f'warning: skipped {x}',file=sys.stderr)
    return sorted(result)

def export(src,root):
    wb=load_workbook(src,data_only=False,keep_vba=src.suffix.lower()=='.xlsm'); count=0
    for ws in wb.worksheets:
        cases=parse(ws)
        if not cases:continue
        dest=root/slug(src.stem)/slug(ws.title);dest.mkdir(parents=True,exist_ok=True);used=set();expected=set()
        for c in cases:
            base=slug(c.name);name=base;n=2
            while name.casefold() in used:name=f'{base}_{n}';n+=1
            used.add(name.casefold());path=dest/f'{name}.md';path.write_text(render(c),encoding='utf-8',newline='\n');expected.add(path.resolve());count+=1
        for old in dest.glob('*.md'):
            if old.resolve() not in expected:old.unlink()
    return count

def main():
    ap=argparse.ArgumentParser();ap.add_argument('inputs',nargs='+');ap.add_argument('-o','--output',default='test-cases-md');a=ap.parse_args()
    files=inputs(a.inputs)
    if not files:ap.error('no input workbooks found')
    total=sum(export(p,Path(a.output)) for p in files);print(f'Exported {total} test case(s) from {len(files)} workbook(s)')
if __name__=='__main__':main()
